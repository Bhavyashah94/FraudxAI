from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

OUTPUT_DIR = Path(r"C:\Users\bhavy\.gemini\antigravity\brain\bb46deb3-54dc-4cb9-a946-6206c822c9a5\scratch\slides")
PPTX_PATH = OUTPUT_DIR / "sample_slide.pptx"

prs = Presentation()
prs.slide_width = Inches(13.333)  # 16:9 widescreen
prs.slide_height = Inches(7.5)

blank_slide_layout = prs.slide_layouts[6]
slide = prs.slides.add_slide(blank_slide_layout)

# Background
bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
bg.fill.solid()
bg.fill.fore_color.rgb = RGBColor(11, 15, 25)
bg.line.fill.background()

# Title Box
txBox = slide.shapes.add_textbox(Inches(0.8), Inches(0.6), Inches(11.7), Inches(1.5))
tf = txBox.text_frame
tf.word_wrap = True

# Tag
p_tag = tf.paragraphs[0]
p_tag.text = "FRAUDXAI • RESEARCH METHODOLOGY & MOTIVATION"
p_tag.font.size = Pt(10)
p_tag.font.bold = True
p_tag.font.color.rgb = RGBColor(56, 189, 248)

# Title
p_title = tf.add_paragraph()
p_title.text = "The Ground-Truth Crisis in Fraud Explainability"
p_title.font.size = Pt(28)
p_title.font.bold = True
p_title.font.color.rgb = RGBColor(255, 255, 255)
p_title.space_before = Pt(6)

# Subtitle
p_sub = tf.add_paragraph()
p_sub.text = "Why post-hoc XAI (SHAP, LIME) cannot be validated on existing fraud datasets, and how causal intervention ground truth resolves the evaluation bottleneck."
p_sub.font.size = Pt(13)
p_sub.font.color.rgb = RGBColor(148, 163, 184)
p_sub.space_before = Pt(4)

# Left Card (Dilemma)
card1 = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(2.2), Inches(5.6), Inches(3.4))
card1.fill.solid()
card1.fill.fore_color.rgb = RGBColor(17, 24, 39)
card1.line.color.rgb = RGBColor(51, 65, 85)

tf1 = card1.text_frame
tf1.word_wrap = True
tf1.margin_left = Inches(0.3)
tf1.margin_right = Inches(0.3)
tf1.margin_top = Inches(0.3)

p_c1_title = tf1.paragraphs[0]
p_c1_title.text = "✖  The Industry Dilemma (Existing Datasets)"
p_c1_title.font.size = Pt(15)
p_c1_title.font.bold = True
p_c1_title.font.color.rgb = RGBColor(244, 63, 94)

bullets1 = [
    "Anonymized PCA Vectors: Public benchmarks (e.g. Kaggle ULB) mask features into eigenvectors (V1–V28), rendering explanations meaningless for banking compliance.",
    "Zero Explanation Ground Truth: Real bank feeds record binary labels (is_fraud), but no mathematical record of the exact root cause behind the alert.",
    "Subjective Confirmation Bias: Prior literature assesses SHAP by visual inspection ('Amount looks high, so SHAP works'), failing rigorous scientific standards."
]

for b in bullets1:
    p = tf1.add_paragraph()
    p.text = "• " + b
    p.font.size = Pt(11)
    p.font.color.rgb = RGBColor(203, 213, 225)
    p.space_before = Pt(8)

# Right Card (Solution)
card2 = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(2.2), Inches(5.7), Inches(3.4))
card2.fill.solid()
card2.fill.fore_color.rgb = RGBColor(17, 24, 39)
card2.line.color.rgb = RGBColor(56, 189, 248)

tf2 = card2.text_frame
tf2.word_wrap = True
tf2.margin_left = Inches(0.3)
tf2.margin_right = Inches(0.3)
tf2.margin_top = Inches(0.3)

p_c2_title = tf2.paragraphs[0]
p_c2_title.text = "✔  Our Causal Intervention Approach"
p_c2_title.font.size = Pt(15)
p_c2_title.font.bold = True
p_c2_title.font.color.rgb = RGBColor(56, 189, 248)

bullets2 = [
    "Physical Attack Footprint (Δx): Every simulated fraud attack injects explicit mathematical interventions against the cardholder's 30-day baseline profile.",
    "Axiomatic Game-Theoretic Ground Truth: Generates exact Shapley attributions via Owen multilinear forms and 128-point Gauss-Legendre path integration.",
    "Standardized Benchmark Metrics: Quantus & OpenXAI evaluation using Precision@k, Recall@k, Kendall's tau, and Relative Attribution Error (RAE)."
]

for b in bullets2:
    p = tf2.add_paragraph()
    p.text = "• " + b
    p.font.size = Pt(11)
    p.font.color.rgb = RGBColor(203, 213, 225)
    p.space_before = Pt(8)

# Metrics Cards
stats = [
    ("62.8%", "INTERVENTION PRECISION (P@3)"),
    ("0.4439", "KENDALL'S TAU CONCORDANCE"),
    ("132 / 132", "INVARIANTS 100% GREEN"),
    ("< 900 km/h", "KINEMATIC TRANSIT CEILING")
]

for i, (val, lbl) in enumerate(stats):
    left = Inches(0.8 + i * 2.95)
    sc = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, Inches(5.8), Inches(2.8), Inches(0.95))
    sc.fill.solid()
    sc.fill.fore_color.rgb = RGBColor(17, 24, 39)
    sc.line.color.rgb = RGBColor(30, 41, 59)
    
    stf = sc.text_frame
    stf.word_wrap = True
    stf.margin_top = Inches(0.12)
    stf.margin_left = Inches(0.15)
    
    pv = stf.paragraphs[0]
    pv.text = val
    pv.font.size = Pt(18)
    pv.font.bold = True
    pv.font.color.rgb = RGBColor(56, 189, 248)
    
    pl = stf.add_paragraph()
    pl.text = lbl
    pl.font.size = Pt(8)
    pl.font.color.rgb = RGBColor(148, 163, 184)
    pl.space_before = Pt(2)

prs.save(str(PPTX_PATH))
print(f"Saved PPTX to {PPTX_PATH}")
